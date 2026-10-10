import io
from datetime import date, datetime

import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from parser_mca import parse_pdf_bytes, compare_visits, sha256_bytes
from storage import Storage


# ============================================================
# CONFIGURACIÓN
# ============================================================

st.set_page_config(
    page_title="Mi Semana MCA",
    page_icon="📅",
    layout="wide",
    initial_sidebar_state="collapsed",
)

st.markdown(
    """
<style>
.block-container{
    padding-top:1rem;
    padding-bottom:4rem;
    max-width:1050px;
}

[data-testid="stMetricValue"]{
    font-size:1.7rem;
}

div[data-testid="stExpander"] details{
    border-radius:14px;
}

.visit-card{
    border:1px solid #e2e8f0;
    border-radius:14px;
    padding:12px 14px;
    margin:8px 0;
    background:white;
}

.new-alert{
    border-left:5px solid #f59e0b;
    background:#fff7ed;
    padding:12px;
    border-radius:10px;
    margin:8px 0;
}

@media(max-width:700px){
    .block-container{
        padding-left:.65rem;
        padding-right:.65rem;
    }

    [data-testid="column"]{
        min-width:0!important;
    }

    button[kind="primary"],
    button[kind="secondary"]{
        min-height:46px;
    }
}
</style>
""",
    unsafe_allow_html=True,
)


DAYS = [
    "Lunes",
    "Martes",
    "Miércoles",
    "Jueves",
    "Viernes",
]


# ============================================================
# REGLAS DE PLANIFICACIÓN
# ============================================================

RULES = {
    "140108961": ["Lunes"],             # IGA Nordique
    "140112886": ["Lunes", "Viernes"],  # Walmart x2
    "140114054": ["Lunes", "Viernes"],  # Super C x2

    "140108970": ["Martes"],            # Maxi Clemenceau
    "140109810": ["Martes"],            # Maxi Estimauville
    "140111340": ["Martes"],            # Provigo Seigneuriale

    "140115222": ["Miércoles"],         # IGA Hêtrière
    "140114989": ["Miércoles"],         # Metro Hêtrière

    "140111635": ["Jueves"],            # Larue
    "140113448": ["Jueves"],            # Charlesbourg
}


# ============================================================
# FUNCIONES
# ============================================================

def week_key(v):
    return f"{v['week_start']}|{v['week_end']}"


def week_label(k):
    start, end = k.split("|")

    ds = datetime.fromisoformat(start).strftime("%d/%m")
    de = datetime.fromisoformat(end).strftime("%d/%m/%Y")

    return f"{ds} – {de}"


def default_day(v):

    seq = RULES.get(v["source_id"], [])

    occurrence = max(
        0,
        int(v.get("occurrence", 1)) - 1
    )

    if occurrence < len(seq):
        return seq[occurrence]

    if seq:
        return seq[-1]

    return "Sin asignar"


# ============================================================
# ACTUALIZAR ESTADO DESDE PDF
# ============================================================

def merge_pdf_into_state(state, parsed):

    current = state.setdefault("visits", {})

    # --------------------------------------------------------
    # IMPORTANTE
    #
    # MyWork puede sacar del PDF visitas que ya fueron hechas.
    #
    # Por eso la aplicación funciona de manera ACUMULATIVA.
    #
    # PDF nuevo:
    # - agrega visitas nuevas
    # - actualiza visitas existentes
    # - conserva visitas anteriores
    #
    # NUNCA elimina una visita solamente porque desapareció
    # del PDF más reciente.
    # --------------------------------------------------------

    for v in parsed:

        old = current.get(
            v["visit_key"],
            {}
        )

        current[v["visit_key"]] = {

            **v,

            # conservar planificación manual
            "day": old.get(
                "day",
                default_day(v)
            ),

            "order": old.get(
                "order",
                v.get("occurrence", 1)
            ),

            # conservar estado trabajado
            "done": old.get(
                "done",
                False
            ),

            # conservar notas
            "notes": old.get(
                "notes",
                ""
            ),

            # toda visita detectada está activa
            "active": True,
        }

    # --------------------------------------------------------
    # NO SE DESACTIVAN VISITAS
    #
    # Si desaparecen del PDF nuevo, se mantienen en la app.
    # --------------------------------------------------------

    return state


# ============================================================
# ESTRUCTURA DE GASTOS
# ============================================================

def ensure_expense_shape(state, wk):

    daily = (
        state
        .setdefault("daily", {})
        .setdefault(wk, {})
    )

    for day in DAYS:

        daily.setdefault(
            day,
            {
                "km": 0.0,
                "parking": 0.0,
                "food": 0.0,
                "tolls": 0.0,
                "other": 0.0,
                "notes": "",
            },
        )

    (
        state
        .setdefault("weekly", {})
        .setdefault(
            wk,
            {
                "tesla_charge": 0.0,
                "phone": 0.0,
                "other": 0.0,
            },
        )
    )


# ============================================================
# LEER PDF
# ============================================================

@st.cache_data(show_spinner=False)
def parse_cached(data: bytes):

    return parse_pdf_bytes(data)


# ============================================================
# CREAR EXCEL
# ============================================================

def make_excel(state, wk):

    wb = Workbook()

    # --------------------------------------------------------
    # CALENDARIO
    # --------------------------------------------------------

    ws = wb.active
    ws.title = "Calendario"

    headers = [
        "Día",
        "Local",
        "Dirección",
        "Horas",
        "Estado MCA",
        "Realizada",
        "Notas",
    ]

    ws.append(headers)

    visits = [
        v
        for v in state["visits"].values()
        if (
            week_key(v) == wk
            and v.get("active", True)
        )
    ]

    order_map = {
        d: i
        for i, d in enumerate(
            DAYS + ["Sin asignar"]
        )
    }

    visits.sort(
        key=lambda v: (
            order_map.get(
                v.get(
                    "day",
                    "Sin asignar"
                ),
                99
            ),
            v.get("order", 0),
            v["store"],
        )
    )

    for v in visits:

        ws.append(
            [
                v.get("day"),
                v["store"],
                v["address"],
                v["hours"],
                v["status"],
                "Sí" if v.get("done") else "No",
                v.get("notes", ""),
            ]
        )

    for cell in ws[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

    widths = [
        14,
        34,
        36,
        10,
        18,
        12,
        28,
    ]

    for i, width in enumerate(
        widths,
        1
    ):
        ws.column_dimensions[
            get_column_letter(i)
        ].width = width

    # --------------------------------------------------------
    # KM Y GASTOS
    # --------------------------------------------------------

    ex = wb.create_sheet(
        "Km y gastos diarios"
    )

    ex.append(
        [
            "Día",
            "Km",
            "Estacionamiento",
            "Comida",
            "Peajes",
            "Otros",
            "Notas",
        ]
    )

    for day in DAYS:

        x = state["daily"][wk][day]

        ex.append(
            [
                day,
                x["km"],
                x["parking"],
                x["food"],
                x["tolls"],
                x["other"],
                x.get(
                    "notes",
                    ""
                ),
            ]
        )

    for cell in ex[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="4472C4"
        )

    widths = [
        14,
        10,
        16,
        12,
        12,
        12,
        28,
    ]

    for i, width in enumerate(
        widths,
        1
    ):
        ex.column_dimensions[
            get_column_letter(i)
        ].width = width

    # --------------------------------------------------------
    # RESUMEN
    # --------------------------------------------------------

    sm = wb.create_sheet(
        "Resumen"
    )

    total_hours = sum(
        v["hours"]
        for v in visits
    )

    done_hours = sum(
        v["hours"]
        for v in visits
        if v.get("done")
    )

    km = sum(
        state["daily"][wk][d]["km"]
        for d in DAYS
    )

    daily_exp = sum(

        sum(
            state["daily"][wk][d][k]
            for k in [
                "parking",
                "food",
                "tolls",
                "other",
            ]
        )

        for d in DAYS
    )

    weekly_exp = sum(

        state["weekly"][wk].get(
            k,
            0
        )

        for k in [
            "tesla_charge",
            "phone",
            "other",
        ]
    )

    rate = float(
        state
        .get(
            "settings",
            {}
        )
        .get(
            "hourly_rate",
            21
        )
    )

    tax = float(
        state
        .get(
            "settings",
            {}
        )
        .get(
            "tax_reserve_pct",
            30
        )
    )

    income = (
        done_hours
        * rate
    )

    rows = [

        (
            "Semana",
            week_label(wk)
        ),

        (
            "Horas programadas",
            total_hours
        ),

        (
            "Horas realizadas",
            done_hours
        ),

        (
            "Ingreso estimado realizado",
            income
        ),

        (
            "Km",
            km
        ),

        (
            "Gastos diarios",
            daily_exp
        ),

        (
            "Carga Tesla semanal",
            state["weekly"][wk]
            .get(
                "tesla_charge",
                0
            )
        ),

        (
            "Gastos semanales",
            weekly_exp
        ),

        (
            "Gastos totales",
            daily_exp
            + weekly_exp
        ),

        (
            "Reserva impuestos",
            income
            * tax
            / 100
        ),
    ]

    for row in rows:

        sm.append(row)

    sm.column_dimensions[
        "A"
    ].width = 28

    sm.column_dimensions[
        "B"
    ].width = 22

    output = io.BytesIO()

    wb.save(output)

    output.seek(0)

    return output.getvalue()


# ============================================================
# ALMACENAMIENTO
# ============================================================

store = Storage()

if "state" not in st.session_state:

    st.session_state.state = (
        store.load_state()
    )

state = st.session_state.state


# ============================================================
# CABECERA
# ============================================================

st.title(
    "📅 Mi Semana MCA"
)

st.caption(
    "Carga manual del PDF de MyWork · "
    "comparación automática · "
    "sincronización entre celular y PC"
)


if store.mode == "google_sheets":

    st.success(
        "☁️ Datos sincronizados en Google Sheets",
        icon="✅",
    )

elif store.mode == "supabase":

    st.success(
        "☁️ Sincronización activa con Supabase",
        icon="✅",
    )

else:

    st.info(
        "Modo local de prueba. "
        "Para compartir los mismos datos entre celular y PC "
        "configura Google Sheets o Supabase.",
        icon="💾",
    )


# ============================================================
# CARGAR PDF
# ============================================================

with st.container(
    border=True
):

    st.subheader(
        "🔄 Actualizar desde MyWork"
    )

    up = st.file_uploader(
        "Sube el PDF actualizado de MyWork",
        type=["pdf"],
        label_visibility="collapsed",
    )

    if up is not None:

        data = up.getvalue()

        digest = sha256_bytes(
            data
        )

        try:

            parsed = parse_cached(
                data
            )

        except Exception as e:

            st.error(
                f"No pude leer este PDF: {e}"
            )

            parsed = []

        if parsed:

            latest = (
                store.latest_snapshot()
            )

            # ==================================================
            # PDF YA CARGADO
            # ==================================================

            if (
                latest
                and latest.get(
                    "sha256"
                )
                == digest
            ):

                st.info(
                    "Este PDF ya fue cargado anteriormente."
                )

                if st.button(
                    "♻️ Reaplicar este PDF y reparar semanas",
                    type="primary",
                    use_container_width=True,
                ):

                    state = (
                        merge_pdf_into_state(
                            state,
                            parsed
                        )
                    )

                    for v in parsed:

                        ensure_expense_shape(
                            state,
                            week_key(v)
                        )

                    store.save_state(
                        state
                    )

                    st.session_state.state = state

                    st.success(
                        "PDF reaplicado correctamente. "
                        "Las visitas fueron restauradas sin borrar "
                        "checks, días, notas ni gastos."
                    )

                    st.rerun()

            # ==================================================
            # PDF NUEVO
            # ==================================================

            else:

                old = (
                    (latest or {})
                    .get(
                        "parsed",
                        []
                    )
                    if latest
                    else []
                )

                changes = (
                    compare_visits(
                        old,
                        parsed
                    )
                )

                added = len(
                    changes["added"]
                )

                removed = len(
                    changes["removed"]
                )

                changed = len(
                    changes["changed"]
                )

                if latest:

                    if (
                        added
                        or removed
                        or changed
                    ):

                        st.warning(
                            f"⚠️ Cambios detectados: "
                            f"{added} visita(s) nueva(s), "
                            f"{removed} ya no aparece(n) en este PDF, "
                            f"{changed} modificada(s)."
                        )

                    else:

                        st.success(
                            "El PDF cambió, pero no hay diferencias "
                            "en visitas, horas o estados."
                        )

                else:

                    st.info(
                        f"Primera carga: "
                        f"{len(parsed)} visitas detectadas."
                    )

                # ------------------------------------------------
                # NUEVAS
                # ------------------------------------------------

                if changes["added"]:

                    st.markdown(
                        "**Nuevas visitas**"
                    )

                    for v in changes["added"]:

                        st.markdown(

                            f"""
<div class='new-alert'>

<b>➕ {v['store']}</b><br>

{v['address']}<br>

<b>{v['hours']:.1f} h</b>

· {week_label(week_key(v))}

</div>
""",

                            unsafe_allow_html=True,
                        )

                # ------------------------------------------------
                # MODIFICADAS
                # ------------------------------------------------

                if changes["changed"]:

                    with st.expander(
                        f"Ver {len(changes['changed'])} "
                        f"cambios de datos/estado"
                    ):

                        for x in changes["changed"]:

                            v = x["visit"]

                            parts = []

                            for field, diff in x[
                                "diffs"
                            ].items():

                                parts.append(
                                    f"{field}: "
                                    f"{diff['old']} "
                                    f"→ "
                                    f"{diff['new']}"
                                )

                            st.write(
                                f"• {v['store']} "
                                f"({week_label(week_key(v))}): "
                                + "; ".join(parts)
                            )

                # ------------------------------------------------
                # YA NO APARECEN
                # ------------------------------------------------

                if changes["removed"]:

                    with st.expander(
                        f"Ver {len(changes['removed'])} visita(s) "
                        f"que ya no aparecen en este PDF "
                        f"(se conservarán en la app)"
                    ):

                        for v in changes["removed"]:

                            st.write(
                                f"• {v['store']} — "
                                f"{v['address']} — "
                                f"{week_label(week_key(v))}"
                            )

                # ------------------------------------------------
                # APLICAR
                # ------------------------------------------------

                if st.button(
                    "Aplicar este PDF y guardar cambios",
                    type="primary",
                    use_container_width=True,
                ):

                    state = (
                        merge_pdf_into_state(
                            state,
                            parsed
                        )
                    )

                    for v in parsed:

                        ensure_expense_shape(
                            state,
                            week_key(v)
                        )

                    store.save_state(
                        state
                    )

                    store.save_snapshot(
                        up.name,
                        digest,
                        parsed
                    )

                    st.session_state.state = state

                    st.success(
                        "PDF aplicado. "
                        "Las visitas anteriores fueron conservadas."
                    )

                    st.rerun()


# ============================================================
# VISITAS ACTIVAS
# ============================================================

active_visits = [

    v

    for v in (
        state
        .get(
            "visits",
            {}
        )
        .values()
    )

    if v.get(
        "active",
        True
    )
]


weeks = sorted(
    {
        week_key(v)
        for v in active_visits
    }
)


if not weeks:

    st.info(
        "Carga tu primer PDF de MyWork para comenzar."
    )

    st.stop()


# ============================================================
# SEMANA ACTUAL PRIMERO
# ============================================================

iso_today = (
    date
    .today()
    .isoformat()
)


def score_week(k):

    start, end = k.split("|")

    if (
        start
        <= iso_today
        <= end
    ):
        return 0

    return 1


weeks = sorted(
    weeks,
    key=lambda k: (
        score_week(k),
        k
    ),
)


selected = st.selectbox(
    "Semana",
    weeks,
    format_func=week_label,
    index=0,
)


ensure_expense_shape(
    state,
    selected
)


week_vis = [

    v

    for v in (
        state["visits"]
        .values()
    )

    if (
        week_key(v)
        == selected
        and v.get(
            "active",
            True
        )
    )
]


# ============================================================
# KPI SEMANA
# ============================================================

hours = sum(
    v["hours"]
    for v in week_vis
)

done_h = sum(
    v["hours"]
    for v in week_vis
    if v.get("done")
)


c1, c2, c3, c4 = (
    st.columns(4)
)


c1.metric(
    "Horas",
    f"{hours:.1f}"
)

c2.metric(
    "Realizadas",
    f"{done_h:.1f}"
)

c3.metric(
    "Visitas",
    len(week_vis)
)

c4.metric(
    "Listas",
    sum(
        1
        for v in week_vis
        if v.get("done")
    )
)


# ============================================================
# PESTAÑAS PRINCIPALES
# ============================================================

(
    plan_tab,
    expense_tab,
    summary_tab,
    history_tab,
) = st.tabs(
    [
        "📅 Semana",
        "🚗 Km y gastos",
        "📊 Resumen",
        "🕘 Historial",
    ]
)


# ============================================================
# PLANIFICACIÓN POR DÍA
# ============================================================

with plan_tab:

    day_options = (
        DAYS
        + ["Sin asignar"]
    )

    unassigned = [

        v

        for v in week_vis

        if (
            v.get(
                "day",
                "Sin asignar"
            )
            == "Sin asignar"
        )
    ]


    visible_days = (
        DAYS
        + (
            ["Sin asignar"]
            if unassigned
            else []
        )
    )


    day_tabs = st.tabs(
        [
            (
                "📌 "
                if day
                == "Sin asignar"
                else ""
            )
            + day

            for day
            in visible_days
        ]
    )


    for (
        day,
        day_tab
    ) in zip(
        visible_days,
        day_tabs
    ):

        with day_tab:

            items = [

                v

                for v in week_vis

                if (
                    v.get(
                        "day",
                        "Sin asignar"
                    )
                    == day
                )
            ]


            items.sort(
                key=lambda v: (
                    v.get(
                        "order",
                        0
                    ),
                    v["store"],
                )
            )


            total = sum(
                v["hours"]
                for v in items
            )


            done_day = sum(
                v["hours"]
                for v in items
                if v.get("done")
            )


            k1, k2, k3 = (
                st.columns(3)
            )


            k1.metric(
                "Horas del día",
                f"{total:.1f}"
            )

            k2.metric(
                "Realizadas",
                f"{done_day:.1f}"
            )

            k3.metric(
                "Visitas",
                len(items)
            )


            if not items:

                st.info(
                    "No hay visitas asignadas a este día."
                )

                continue


            for v in items:

                with st.container(
                    border=True
                ):

                    left, right = (
                        st.columns(
                            [4, 1]
                        )
                    )


                    with left:

                        done = (
                            st.checkbox(
                                f"{v['store']} — "
                                f"{v['hours']:.1f} h",
                                value=bool(
                                    v.get(
                                        "done"
                                    )
                                ),
                                key=(
                                    "done_"
                                    + v[
                                        "visit_key"
                                    ]
                                ),
                            )
                        )


                        st.caption(
                            v["address"]
                        )


                        st.caption(
                            f"Estado MCA: "
                            f"{v['status']}"
                        )


                    with right:

                        current_day = (
                            v.get(
                                "day",
                                "Sin asignar"
                            )
                        )

                        if (
                            current_day
                            not in
                            day_options
                        ):

                            current_day = (
                                "Sin asignar"
                            )


                        new_day = (
                            st.selectbox(
                                "Mover a",
                                day_options,
                                index=(
                                    day_options
                                    .index(
                                        current_day
                                    )
                                ),
                                key=(
                                    "day_"
                                    + v[
                                        "visit_key"
                                    ]
                                ),
                            )
                        )


                    notes = (
                        st.text_input(
                            "Nota",
                            value=v.get(
                                "notes",
                                ""
                            ),
                            key=(
                                "note_"
                                + v[
                                    "visit_key"
                                ]
                            ),
                            placeholder=(
                                "Nota opcional"
                            ),
                        )
                    )


                    if (
                        done
                        != v.get(
                            "done"
                        )
                        or
                        new_day
                        != v.get(
                            "day"
                        )
                        or
                        notes
                        != v.get(
                            "notes",
                            ""
                        )
                    ):

                        v["done"] = done
                        v["day"] = new_day
                        v["notes"] = notes

                        store.save_state(
                            state
                        )


# ============================================================
# KM Y GASTOS
# ============================================================

with expense_tab:

    st.subheader(
        "Registro diario"
    )


    day = st.selectbox(
        "Día a registrar",
        DAYS,
        key="expense_day"
    )


    x = (
        state["daily"]
        [selected]
        [day]
    )


    a, b = st.columns(2)


    km = a.number_input(
        "Kilómetros",
        min_value=0.0,
        value=float(
            x["km"]
        ),
        step=1.0,
    )


    parking = b.number_input(
        "Estacionamiento ($)",
        min_value=0.0,
        value=float(
            x["parking"]
        ),
        step=1.0,
    )


    c, d = st.columns(2)


    food = c.number_input(
        "Comida ($)",
        min_value=0.0,
        value=float(
            x["food"]
        ),
        step=1.0,
    )


    tolls = d.number_input(
        "Peajes ($)",
        min_value=0.0,
        value=float(
            x["tolls"]
        ),
        step=1.0,
    )


    other = st.number_input(
        "Otros gastos del día ($)",
        min_value=0.0,
        value=float(
            x["other"]
        ),
        step=1.0,
    )


    notes = st.text_input(
        "Nota del día",
        value=x.get(
            "notes",
            ""
        ),
    )


    if st.button(
        "Guardar día",
        type="primary",
        use_container_width=True,
    ):

        state["daily"][
            selected
        ][day] = {

            "km": km,

            "parking": parking,

            "food": food,

            "tolls": tolls,

            "other": other,

            "notes": notes,
        }


        store.save_state(
            state
        )


        st.success(
            "Registro diario guardado."
        )


    st.divider()


    st.subheader(
        "Gastos semanales"
    )


    weekly = (
        state["weekly"]
        [selected]
    )


    tesla = st.number_input(
        "Carga/energía Tesla de la semana ($)",
        min_value=0.0,
        value=float(
            weekly.get(
                "tesla_charge",
                0
            )
        ),
        step=1.0,
    )


    phone = st.number_input(
        "Teléfono/datos atribuible a esta semana ($)",
        min_value=0.0,
        value=float(
            weekly.get(
                "phone",
                0
            )
        ),
        step=1.0,
    )


    weekly_other = (
        st.number_input(
            "Otros gastos semanales ($)",
            min_value=0.0,
            value=float(
                weekly.get(
                    "other",
                    0
                )
            ),
            step=1.0,
        )
    )


    if st.button(
        "Guardar gastos semanales",
        use_container_width=True,
    ):

        weekly.update(
            {
                "tesla_charge": tesla,
                "phone": phone,
                "other": weekly_other,
            }
        )


        store.save_state(
            state
        )


        st.success(
            "Gastos semanales guardados."
        )


# ============================================================
# RESUMEN
# ============================================================

with summary_tab:

    settings = (
        state.setdefault(
            "settings",
            {
                "hourly_rate": 21.0,
                "tax_reserve_pct": 30.0,
            },
        )
    )


    ca, cb = st.columns(2)


    rate = ca.number_input(
        "Tarifa por hora ($)",
        min_value=0.0,
        value=float(
            settings.get(
                "hourly_rate",
                21
            )
        ),
        step=0.5,
    )


    tax = cb.number_input(
        "Reserva para impuestos (%)",
        min_value=0.0,
        max_value=100.0,
        value=float(
            settings.get(
                "tax_reserve_pct",
                30
            )
        ),
        step=1.0,
    )


    settings[
        "hourly_rate"
    ] = rate


    settings[
        "tax_reserve_pct"
    ] = tax


    store.save_state(
        state
    )


    km_total = sum(
        state["daily"]
        [selected]
        [d]["km"]
        for d in DAYS
    )


    daily_exp = sum(

        sum(
            state["daily"]
            [selected]
            [d][k]

            for k in [
                "parking",
                "food",
                "tolls",
                "other",
            ]
        )

        for d in DAYS
    )


    weekly_exp = sum(

        state["weekly"]
        [selected]
        .get(
            k,
            0
        )

        for k in [
            "tesla_charge",
            "phone",
            "other",
        ]
    )


    income = (
        done_h
        * rate
    )


    r1, r2, r3, r4 = (
        st.columns(4)
    )


    r1.metric(
        "Ingreso realizado",
        f"${income:,.2f}"
    )


    r2.metric(
        "Km semana",
        f"{km_total:.0f}"
    )


    r3.metric(
        "Gastos registrados",
        f"${daily_exp + weekly_exp:,.2f}"
    )


    r4.metric(
        "Reserva impuestos",
        f"${income * tax / 100:,.2f}"
    )


    st.caption(
        "La reserva para impuestos es una provisión personal, "
        "no una determinación fiscal de deducibilidad."
    )


    xlsx = make_excel(
        state,
        selected
    )


    st.download_button(
        "⬇️ Descargar Excel de esta semana",
        data=xlsx,
        file_name=(
            f"Mi_Semana_MCA_"
            f"{selected.split('|')[0]}"
            f".xlsx"
        ),
        mime=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        use_container_width=True,
    )


    if (
        store.mode
        == "google_sheets"
    ):

        st.success(
            "Google Sheets está funcionando como almacenamiento central. "
            "Lo que guardes en el celular queda disponible también en el PC."
        )

        st.caption(
            "El Excel es una copia descargable de la semana seleccionada."
        )

    else:

        st.caption(
            "El Excel se genera con la información actual. "
            "Si quieres usar Google Sheets como almacenamiento central, "
            "configura GOOGLE_SHEET_ID y la cuenta de servicio "
            "en Streamlit Secrets."
        )


# ============================================================
# HISTORIAL
# ============================================================

with history_tab:

    hist = store.history(
        20
    )


    if not hist:

        st.caption(
            "Aún no hay historial de PDFs."
        )


    for h in hist:

        created = h.get(
            "created_at",
            ""
        )


        try:

            created = (
                datetime
                .fromisoformat(
                    created.replace(
                        "Z",
                        "+00:00"
                    )
                )
                .strftime(
                    "%d/%m/%Y %H:%M"
                )
            )

        except Exception:

            pass


        st.write(
            f"**{created}** · "
            f"{h.get('filename', 'PDF')} · "
            f"{len(h.get('parsed', []))} visitas"
        )
